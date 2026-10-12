import io
import unittest

from docx import Document

from src import config, export, mdparse, render
from tests.fakes import base_exemplo, limpar
from tests.test_render import gerado_bom

BR = "  "+chr(10)   # quebra dura do Markdown: dois espaços no fim da linha
NL = chr(10)
MD = NL.join([
    "# Mariana Exemplo" + BR[:-1],
    "**Analista de Marketing**" + BR[:-1],
    "São Paulo, Brasil" + BR[:-1],
    "+55 11 90000-0000 | mariana@example.com | linkedin.com/in/mariana", "",
    "## Resumo", "", "Analista com **foco** em _mídia paga_ e [portfólio](https://exemplo.com/x).", "",
    "## Experiência Profissional", "", "**Empresa Exemplo**", "",
    "_Analista · Março 2022 – Presente · São Paulo_", "",
    "* Reduziu o custo por lead em 18%.", "- Estruturou o funil de 12 mil usuários.", ""])


class TestMdparse(unittest.TestCase):
    def test_blocos_e_quebras(self):
        b = mdparse.parse(MD)
        self.assertEqual([x["tipo"] for x in b], ["h1", "p", "h2", "p", "h2", "p", "p", "li", "li"])
        self.assertEqual(len(b[1]["linhas"]), 3)   # título, local, contatos: quebras duras
        self.assertEqual(mdparse.texto_simples(b[0]["linhas"][0]), "Mariana Exemplo")

    def test_inline(self):
        t = mdparse.parse_inline("Com **foco** em _mídia_ e [site](https://a.com/x) e mariana_x@a.com")
        self.assertIn(("foco", frozenset({"b"})), t)
        self.assertIn(("mídia", frozenset({"i"})), t)
        simples = mdparse.texto_simples(t)
        self.assertIn("site (https://a.com/x)", simples)
        self.assertIn("mariana_x@a.com", simples)   # underscore dentro de palavra não vira itálico

    def test_empresa_e_cargo(self):
        b = mdparse.parse(MD)
        self.assertTrue(mdparse.so_estilo(b[5]["linhas"][0], "b"))
        self.assertTrue(mdparse.so_estilo(b[6]["linhas"][0], "i"))

    def test_quebra_suave_junta_linhas(self):
        b = mdparse.parse("linha um\nlinha dois\n\noutro")
        self.assertEqual(mdparse.texto_simples(b[0]["linhas"][0]), "linha um linha dois")


class TestExport(unittest.TestCase):
    def setUp(self):
        self.kb, self.tmp = base_exemplo()
        cv = render.montar(self.kb, gerado_bom(), "pt", config.carregar_config())
        self.md = cv.markdown

    def tearDown(self):
        limpar(self.tmp)

    def test_pdf_tem_nome_secoes_e_texto_selecionavel(self):
        dados, paginas = export.gerar_pdf(self.md)
        self.assertTrue(dados.startswith(b"%PDF"))
        self.assertEqual(paginas, 1)
        texto = export.texto_do_pdf(dados)
        for esperado in ("Mariana Exemplo", "Resumo", "Competências", "Experiência Profissional", "Formação",
                         "Empresa Exemplo Digital", "Março 2022", "linkedin.com/in/mariana-exemplo",
                         "Estruturou o funil", "90000-0000"):
            self.assertIn(esperado, texto)
        self.assertNotIn("**", texto)
        self.assertNotIn("](", texto)

    def test_pdf_caracteres_fora_do_cp1252_nao_derrubam(self):
        md = self.md + "\n* Seta → emoji \U0001F680 check ✓ aspas “x” reticências… chinês 中\n"
        dados, _ = export.gerar_pdf(md)
        texto = export.texto_do_pdf(dados)
        self.assertIn("Seta -> emoji", texto)
        self.assertNotIn("\U0001F680", texto)

    def test_sanear(self):
        self.assertEqual(export.sanear_cp1252("a→b c​d"), "a->b cd")
        self.assertEqual(export.sanear_cp1252("ação – ok · •"), "ação – ok · •")
        self.assertEqual(export.sanear_cp1252("\U0001F680"), "")

    def test_pdf_conta_paginas(self):
        longo = self.md + "\n" + "\n".join(f"* Bullet numero {i} com texto para ocupar espaço na página inteira ok." for i in range(120))
        _, paginas = export.gerar_pdf(longo)
        self.assertGreaterEqual(paginas, 2)

    def test_pdf_markdown_vazio(self):
        dados, paginas = export.gerar_pdf("")
        self.assertTrue(dados.startswith(b"%PDF"))
        self.assertEqual(paginas, 1)

    def test_docx_estilos_nativos_e_conteudo(self):
        dados = export.gerar_docx(self.md)
        doc = Document(io.BytesIO(dados))
        estilos = [p.style.name for p in doc.paragraphs]
        self.assertEqual(estilos[0], "Heading 1")
        self.assertIn("Heading 2", estilos)
        self.assertIn("List Bullet", estilos)
        texto = "\n".join(p.text for p in doc.paragraphs)
        for esperado in ("Mariana Exemplo", "Experiência Profissional", "Estruturou o funil", "linkedin.com/in/mariana-exemplo"):
            self.assertIn(esperado, texto)
        self.assertNotIn("**", texto)
        self.assertEqual(len(doc.tables), 0)

    def test_docx_sem_fonte_de_tema_nos_titulos(self):
        doc = Document(io.BytesIO(export.gerar_docx(self.md)))
        from docx.oxml.ns import qn
        for nome in ("Heading 1", "Heading 2"):
            fontes = doc.styles[nome].element.rPr.find(qn("w:rFonts"))
            self.assertEqual(fontes.get(qn("w:ascii")), "Arial")
            self.assertFalse([a for a in fontes.attrib if "Theme" in a])

    def test_docx_aguenta_texto_estranho(self):
        export.gerar_docx(self.md + "\n* controle \x00\x01 e emoji \U0001F680\n")


if __name__ == "__main__":
    unittest.main()
