"""Utilidades de texto: idioma, JSON tolerante, leitura de arquivos, máscara de contato."""
from __future__ import annotations

import io
import json
import re
import unicodedata

# ---------------------------------------------------------------- normalização

def normalizar(texto: str) -> str:
    """Minúsculas, sem acento, espaços colapsados. Para comparar textos."""
    t = unicodedata.normalize("NFKD", texto or "")
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"\s+", " ", t.lower()).strip()


def slugify(texto: str) -> str:
    t = normalizar(texto)
    t = re.sub(r"[^a-z0-9]+", "-", t).strip("-")
    return t or "sem-nome"


# ---------------------------------------------------------------------- idioma

_STOP_PT = set("""de da do das dos para com uma uns umas nao voce vocês ser serao sera por ao na no nas nos
que como mais ou sua seu suas seus tambem sobre entre atuar experiencia equipe empresa vaga requisitos
desejavel obrigatorio conhecimento habilidades responsabilidades area""".split())
_STOP_EN = set("""the and of to with for in you your our we is are will be this that from as on at or
experience team company role requirements preferred required knowledge skills responsibilities about
who what have has an by their they""".split())


def detectar_idioma(texto: str, min_palavras: int = 15) -> str | None:
    """'pt', 'en' ou None quando a heurística não tem segurança."""
    palavras = re.findall(r"[a-zA-ZÀ-ÿ]+", normalizar(texto))
    if len(palavras) < min_palavras:
        return None
    pt = sum(1 for p in palavras if p in _STOP_PT)
    en = sum(1 for p in palavras if p in _STOP_EN)
    if pt == en or max(pt, en) < 3:
        return None
    return "pt" if pt > en else "en"


# ------------------------------------------------------------------------ JSON

_THINK = re.compile(r"<think>.*?</think>", re.S | re.I)
_VIRGULA_SOBRANDO = re.compile(r",\s*([}\]])")


def remover_raciocinio(texto: str) -> str:
    return _THINK.sub("", texto or "").strip()


def extrair_json(texto: str):
    """Extrai o primeiro objeto/lista JSON de uma resposta de modelo.

    Tolera cerca ```json, texto antes e depois e vírgula sobrando.
    Levanta ValueError se não achar JSON válido.
    """
    t = remover_raciocinio(texto)
    cerca = re.search(r"```(?:json)?\s*(.*?)```", t, re.S | re.I)
    candidatos = [cerca.group(1)] if cerca else []
    candidatos.append(t)
    decoder = json.JSONDecoder()
    for cand in candidatos:
        for variante in (cand, _VIRGULA_SOBRANDO.sub(r"\1", cand)):
            for m in re.finditer(r"[{\[]", variante):
                try:
                    obj, _ = decoder.raw_decode(variante[m.start():])
                except ValueError:
                    continue
                if isinstance(obj, (dict, list)):
                    return obj
    raise ValueError("Nenhum JSON válido encontrado na resposta.")


# ------------------------------------------------------------ máscara de contato

_EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
_URL = re.compile(
    r"(?:https?://|www\.)\S+|\b(?:[\w-]+\.)?(?:linkedin|github|gitlab|behance|instagram|facebook)\.com\S*",
    re.I,
)
# Telefones: com DDI (+55 ...), com DDD entre parênteses, ou 5-4 dígitos com hífen.
# Propositalmente NÃO casa "2016 2021", "2016-2021" nem "2.000".
_TELEFONES = [
    re.compile(r"\+\d{1,3}[\s.-]?\(?\d{1,3}\)?[\s.-]?\d{3,5}[\s.-]?\d{4}\b"),
    re.compile(r"\(\d{2,3}\)\s?\d{4,5}[\s.-]?\d{4}\b"),
    re.compile(r"\b\d{2}\s9\d{4}[\s.-]?\d{4}\b"),
    re.compile(r"\b\d{5}-\d{4}\b"),
]


def mascarar_contato(texto: str) -> str:
    """Troca e-mail, URL e telefone por marcadores. Aplicada a tudo que vai ao modelo."""
    t = _EMAIL.sub("[e-mail omitido]", texto or "")
    t = _URL.sub("[link omitido]", t)
    for padrao in _TELEFONES:
        t = padrao.sub("[telefone omitido]", t)
    return t


def contem_contato(texto: str) -> bool:
    return mascarar_contato(texto) != (texto or "")


def extrair_contato_local(texto: str) -> dict:
    """Acha e-mail, telefone e LinkedIn com regex, sem passar pelo modelo."""
    achado: dict = {}
    m = _EMAIL.search(texto)
    if m:
        achado["email"] = m.group(0)
    m = re.search(r"linkedin\.com/in/[\w\-%]+", texto, re.I)
    if m:
        achado["linkedin"] = m.group(0)
    for padrao in _TELEFONES[:3]:
        m = padrao.search(texto)
        if m:
            achado["telefone"] = m.group(0).strip()
            break
    return achado


_MESES = {
    "janeiro": 1, "jan": 1, "january": 1, "fevereiro": 2, "fev": 2, "february": 2, "feb": 2, "marco": 3, "mar": 3,
    "march": 3, "abril": 4, "abr": 4, "april": 4, "apr": 4, "maio": 5, "mai": 5, "may": 5, "junho": 6, "jun": 6,
    "june": 6, "julho": 7, "jul": 7, "july": 7, "agosto": 8, "ago": 8, "august": 8, "aug": 8, "setembro": 9,
    "set": 9, "sep": 9, "sept": 9, "september": 9, "outubro": 10, "out": 10, "october": 10, "oct": 10,
    "novembro": 11, "nov": 11, "november": 11, "dezembro": 12, "dez": 12, "december": 12, "dec": 12,
}
_MES_ANO = re.compile(r"\b([a-z]{3,9})\.?\s*(?:de\s+|del\s+|/|-|,)?\s*((?:19|20)\d\d)\b")
_NUM_ANO = re.compile(r"\b(0?[1-9]|1[0-2])\s*[/.-]\s*((?:19|20)\d\d)\b")
_ANO_NUM = re.compile(r"\b((?:19|20)\d\d)\s*[/.-]\s*(0?[1-9]|1[0-2])\b")


def mes_citado(material: str, aaaa_mm: str) -> bool:
    """True se o mês e o ano de `aaaa_mm` aparecem juntos no material ("março de 2022", "03/2022", "2022-03").

    Serve para não aceitar um mês que o modelo completou sozinho quando a pessoa só disse o ano.
    """
    ano, mes = aaaa_mm.split("-")
    ano, mes = int(ano), int(mes)
    t = normalizar(material)
    for m in _MES_ANO.finditer(t):
        if _MESES.get(m.group(1)) == mes and int(m.group(2)) == ano:
            return True
    for m in _NUM_ANO.finditer(t):
        if int(m.group(1)) == mes and int(m.group(2)) == ano:
            return True
    for m in _ANO_NUM.finditer(t):
        if int(m.group(1)) == ano and int(m.group(2)) == mes:
            return True
    return False


def palpite_nome(texto: str) -> str:
    """Primeira linha do currículo que parece um nome (2 a 5 palavras, sem dígito, e-mail ou símbolo)."""
    for linha in (texto or "").splitlines()[:8]:
        l = linha.strip()
        palavras = l.split()
        if 2 <= len(palavras) <= 5 and not re.search(r"[\d@:/|•·]", l) and all(p[:1].isupper() or p.lower() in
                                                                           ("da", "de", "do", "das", "dos", "e") for p in palavras):
            return l
    return ""


# ------------------------------------------------------------ leitura de arquivos

def ler_arquivo(nome: str, dados: bytes) -> str:
    """Texto de PDF, DOCX, MD ou TXT enviado pelo usuário."""
    ext = nome.lower().rsplit(".", 1)[-1] if "." in nome else ""
    if ext == "pdf":
        from pypdf import PdfReader
        leitor = PdfReader(io.BytesIO(dados))
        return "\n".join((p.extract_text() or "") for p in leitor.pages).strip()
    if ext == "docx":
        from docx import Document
        doc = Document(io.BytesIO(dados))
        partes = [p.text for p in doc.paragraphs]
        for tabela in doc.tables:
            for linha in tabela.rows:
                partes.append(" | ".join(c.text for c in linha.cells))
        return "\n".join(partes).strip()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return dados.decode(enc).strip()
        except UnicodeDecodeError:
            continue
    return dados.decode("utf-8", errors="replace").strip()
