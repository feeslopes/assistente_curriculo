"""Exportação ATS: PDF (reportlab) e DOCX (python-docx), ambos a partir do mesmo Markdown.

PDF: coluna única, A4, fonte padrão Helvetica, texto selecionável, sem tabela, imagem, ícone ou caixa de
texto. A Helvetica padrão usa WinAnsi (cp1252): o que não couber nele é trocado antes de desenhar.
DOCX: estilos nativos Heading 1, Heading 2 e List Bullet, que os parsers de ATS reconhecem.
"""
from __future__ import annotations

import io
import unicodedata
from xml.sax.saxutils import escape

from . import mdparse

TROCAS = {
    "→": "->", "←": "<-", "⇒": "=>", "✓": "v", "✔": "v", "•": "•",
    "−": "-", "‐": "-", "‑": "-", "‒": "-", "―": "-", " ": " ", " ": " ",
    " ": " ", "​": "", "‌": "", "‍": "", "﻿": "", "′": "'", "″": '"',
    "⁃": "-", "●": "•", "▪": "•", "‣": "•", "∙": "•",
}


def sanear_cp1252(texto: str) -> str:
    """Troca ou remove o que a Helvetica padrão não desenha (setas, emojis, símbolos), sem derrubar nada."""
    saida = []
    for ch in texto:
        ch = TROCAS.get(ch, ch)
        for c in ch:
            try:
                c.encode("cp1252")
                saida.append(c)
            except UnicodeEncodeError:
                base = unicodedata.normalize("NFKD", c)
                base = "".join(x for x in base if not unicodedata.combining(x))
                try:
                    base.encode("cp1252")
                    saida.append(base)
                except UnicodeEncodeError:
                    pass  # emoji e afins: descarta
    return "".join(saida)


# ---------------------------------------------------------------------------- PDF

def _xml(linha) -> str:
    partes = []
    for texto, est in linha:
        t = escape(sanear_cp1252(texto))
        if "b" in est:
            t = f"<b>{t}</b>"
        if "i" in est:
            t = f"<i>{t}</i>"
        partes.append(t)
    return "".join(partes)


def gerar_pdf(markdown: str) -> tuple[bytes, int]:
    """Devolve (bytes do PDF, número de páginas)."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    base = dict(fontName="Helvetica", fontSize=10, leading=14.2, textColor=colors.HexColor("#111111"), alignment=TA_LEFT)
    est = {
        "h1": ParagraphStyle("h1", **{**base, "fontName": "Helvetica-Bold", "fontSize": 18, "leading": 22, "spaceAfter": 2}),
        "p": ParagraphStyle("p", **{**base, "spaceAfter": 1}),
        "h2": ParagraphStyle("h2", **{**base, "fontName": "Helvetica-Bold", "fontSize": 11.5, "leading": 14,
                                       "spaceBefore": 11, "spaceAfter": 2, "keepWithNext": 1}),
        "h3": ParagraphStyle("h3", **{**base, "fontName": "Helvetica-Bold", "spaceBefore": 4, "keepWithNext": 1}),
        "empresa": ParagraphStyle("empresa", **{**base, "fontName": "Helvetica-Bold", "fontSize": 10.5, "spaceBefore": 5,
                                                 "spaceAfter": 0, "keepWithNext": 1}),
        "cargo": ParagraphStyle("cargo", **{**base, "fontName": "Helvetica-Bold", "textColor": colors.HexColor("#555555"),
                                             "spaceBefore": 4, "spaceAfter": 2, "keepWithNext": 1}),
        "li": ParagraphStyle("li", **{**base, "leftIndent": 13, "bulletIndent": 3, "spaceAfter": 1.8}),
    }

    blocos = mdparse.parse(markdown)
    fluxo: list = []
    nome = ""
    for b in blocos:
        tipo, linhas = b["tipo"], b["linhas"]
        if tipo == "h1":
            nome = nome or mdparse.texto_simples(linhas[0])
            fluxo.append(Paragraph(_xml(linhas[0]), est["h1"]))
        elif tipo == "h2":
            fluxo.append(Paragraph(_xml(linhas[0]), est["h2"]))
            # linha fina desenhada (não é texto, o ATS não lê)
            fluxo.append(HRFlowable(width="100%", thickness=0.6, color=colors.HexColor("#999999"),
                                    spaceBefore=0, spaceAfter=4))
        elif tipo == "h3":
            fluxo.append(Paragraph(_xml(linhas[0]), est["h3"]))
        elif tipo == "li":
            fluxo.append(Paragraph(_xml(linhas[0]), est["li"], bulletText="•"))
        else:
            if len(linhas) == 1 and mdparse.so_estilo(linhas[0], "b"):
                semi = [(t, frozenset(e - {"b"})) for t, e in linhas[0]]
                fluxo.append(Paragraph(f"<b>{_xml(semi)}</b>", est["empresa"]))
            elif len(linhas) == 1 and mdparse.so_estilo(linhas[0], "i"):
                semi = [(t, frozenset(e - {"i"})) for t, e in linhas[0]]
                fluxo.append(Paragraph(_xml(semi), est["cargo"]))
            else:
                fluxo.append(Paragraph("<br/>".join(_xml(l) for l in linhas), est["p"]))
    if not fluxo:
        fluxo.append(Spacer(1, 1))

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=13 * mm,
                            bottomMargin=13 * mm, title=sanear_cp1252(nome) or "Currículo",
                            author=sanear_cp1252(nome), subject="Currículo")
    doc.build(fluxo)
    return buf.getvalue(), doc.page


# --------------------------------------------------------------------------- DOCX

def gerar_docx(markdown: str, fonte: str = "Arial") -> bytes:
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin = sec.bottom_margin = Mm(15)

    def ajustar(estilo, tam, negrito=None, cor=None):
        estilo.font.name = fonte
        estilo.font.size = Pt(tam)
        if negrito is not None:
            estilo.font.bold = negrito
        estilo.font.color.rgb = RGBColor(*(cor or (0, 0, 0)))
        rpr = estilo.element.get_or_add_rPr()
        fontes = rpr.find(qn("w:rFonts"))
        if fontes is None:
            fontes = rpr.makeelement(qn("w:rFonts"), {})
            rpr.append(fontes)
        # a fonte do tema sobrepõe a escolhida se estes atributos ficarem
        for atributo in list(fontes.attrib):
            if "Theme" in atributo:
                del fontes.attrib[atributo]
        for atributo in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            fontes.set(qn(atributo), fonte)

    estilos = doc.styles
    ajustar(estilos["Normal"], 10)
    ajustar(estilos["Heading 1"], 18, True)
    ajustar(estilos["Heading 2"], 11.5, True)
    ajustar(estilos["Heading 3"], 10.5, True)
    ajustar(estilos["List Bullet"], 10)
    for nome, antes, depois in (("Normal", 0, 2), ("Heading 1", 0, 2), ("Heading 2", 10, 3), ("Heading 3", 4, 1),
                                ("List Bullet", 0, 2)):
        pf = estilos[nome].paragraph_format
        pf.space_before, pf.space_after = Pt(antes), Pt(depois)
    estilos["Heading 2"].paragraph_format.keep_with_next = True

    def escrever(par, linha):
        for texto, est in linha:
            run = par.add_run(sanear_docx(texto))
            run.bold = "b" in est or None
            run.italic = "i" in est or None

    def sanear_docx(t: str) -> str:
        return "".join(c for c in t if c == "\t" or c == "\n" or ord(c) >= 32)

    for b in mdparse.parse(markdown):
        tipo, linhas = b["tipo"], b["linhas"]
        if tipo in ("h1", "h2", "h3"):
            p = doc.add_paragraph(style={"h1": "Heading 1", "h2": "Heading 2", "h3": "Heading 3"}[tipo])
            escrever(p, linhas[0])
        elif tipo == "li":
            p = doc.add_paragraph(style="List Bullet")
            escrever(p, linhas[0])
        else:
            p = doc.add_paragraph()
            for i, linha in enumerate(linhas):
                if i:
                    p.add_run().add_break(WD_BREAK.LINE)
                escrever(p, linha)
            if len(linhas) == 1 and mdparse.so_estilo(linhas[0], "b"):
                p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_before = Pt(4)
            elif len(linhas) == 1 and mdparse.so_estilo(linhas[0], "i"):
                p.paragraph_format.keep_with_next = True
                for r in p.runs:
                    r.italic = None
                    r.bold = True
                    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------- utilidades

def texto_do_pdf(dados: bytes) -> str:
    """Extrai o texto como um parser de ATS faria."""
    from pypdf import PdfReader
    return "\n".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(dados)).pages)
